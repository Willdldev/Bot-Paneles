"""
Arma el resumen corto que se manda directo en el chat y el Excel de
detalle que se ofrece a pedido (botón), para no saturar el chat cuando
hay muchas combinaciones de marca/modelo/potencia.
"""
import io

from telegram import InlineKeyboardButton, InlineKeyboardMarkup
from openpyxl import Workbook
from openpyxl.styles import Font

from db import get_pool


async def _filas_stock(conn):
    return await conn.fetch(
        "SELECT marca, modelo, potencia_w, en_almacen, pendiente_por_llegar, reservado, danados "
        "FROM paneles_stock ORDER BY marca, modelo, potencia_w"
    )


async def _reservas_activas(conn, marca, modelo, potencia_w):
    return await conn.fetch(
        "SELECT proyecto, cantidad FROM reservas "
        "WHERE marca=$1 AND modelo=$2 AND potencia_w=$3 AND estado_despacho != 'despachada' "
        "ORDER BY fecha_solicitud",
        marca, modelo, potencia_w,
    )


def _tabla(encabezados: list[str], filas: list[list], max_ancho: int = 18) -> str:
    """
    Arma una tabla en texto monoespaciado (dentro de un bloque ``` para
    que Telegram la muestre con fuente de ancho fijo y las columnas
    queden alineadas). Las columnas donde todos los valores son
    numéricos se alinean a la derecha; el resto, a la izquierda.
    """
    filas_txt = [[str(v) for v in fila] for fila in filas]

    anchos = []
    alinear_derecha = []
    for i, encabezado in enumerate(encabezados):
        valores_col = [f[i] for f in filas_txt]
        ancho = max([len(encabezado)] + [min(len(v), max_ancho) for v in valores_col])
        anchos.append(ancho)
        es_numerica = bool(valores_col) and all(
            v.lstrip("-").isdigit() for v in valores_col
        )
        alinear_derecha.append(es_numerica)

    def _celda(valor: str, ancho: int, derecha: bool) -> str:
        if len(valor) > ancho:
            valor = valor[: ancho - 1] + "…"
        return valor.rjust(ancho) if derecha else valor.ljust(ancho)

    def _fila(valores: list[str]) -> str:
        return " ".join(_celda(v, a, d) for v, a, d in zip(valores, anchos, alinear_derecha))

    lineas = [_fila(encabezados), " ".join("-" * a for a in anchos)]
    lineas += [_fila(f) for f in filas_txt]
    return "```\n" + "\n".join(lineas) + "\n```"


async def construir_reporte(vista: str):
    """
    vista='fisica'    -> lo que hay en la bodega ya mismo (Almacén / Inventario paneles):
                          disponible = en_almacen - min(reservado, en_almacen)
    vista='comercial' -> lo que se puede vender en total (Reservar Paneles):
                          disponible = en_almacen + pendiente_por_llegar - reservado

    Devuelve (resumen_en_texto, filas_para_excel).
    """
    pool = await get_pool()
    filas_reporte = []
    async with pool.acquire() as conn:
        for fila in await _filas_stock(conn):
            marca, modelo, potencia = fila["marca"], fila["modelo"], fila["potencia_w"]
            en_almacen = fila["en_almacen"]
            pendiente = fila["pendiente_por_llegar"]
            reservado_total = fila["reservado"]
            reservas = await _reservas_activas(conn, marca, modelo, potencia)

            if vista == "fisica":
                reservado_mostrado = min(reservado_total, en_almacen)
                disponible = en_almacen - reservado_mostrado
            else:
                reservado_mostrado = reservado_total
                disponible = en_almacen + pendiente - reservado_total

            filas_reporte.append({
                "marca": marca, "modelo": modelo, "potencia_w": potencia,
                "en_almacen": en_almacen, "pendiente_por_llegar": pendiente,
                "disponible": disponible, "reservado": reservado_mostrado,
                "danados": fila["danados"], "proyectos": reservas,
            })

    if not filas_reporte:
        return "No hay paneles registrados todavía.", filas_reporte

    if vista == "fisica":
        titulo = "📦 *Inventario físico*"
        encabezados = ["Marca", "Modelo", "Pot", "Almac", "Disp", "Reserv", "Dañ"]
        filas_tabla = [
            [f["marca"], f["modelo"], f["potencia_w"], f["en_almacen"],
             f["disponible"], f["reservado"], f["danados"]]
            for f in filas_reporte
        ]
    else:
        titulo = "🟢 *Disponible para vender*"
        encabezados = ["Marca", "Modelo", "Pot", "Disp", "Reserv"]
        filas_tabla = [
            [f["marca"], f["modelo"], f["potencia_w"], f["disponible"], f["reservado"]]
            for f in filas_reporte
        ]

    return f"{titulo}\n{_tabla(encabezados, filas_tabla)}", filas_reporte


def construir_excel(filas_reporte, vista: str) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    encabezados = ["Marca", "Modelo", "Potencia (W)", "En almacén"]
    if vista == "comercial":
        encabezados.append("Pendiente por llegar")
    encabezados += ["Disponible", "Reservado", "Proyectos (reservado)"]
    if vista == "fisica":
        encabezados.append("Dañados")
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for f in filas_reporte:
        proyectos_txt = ", ".join(f"{r['proyecto']} ({r['cantidad']})" for r in f["proyectos"]) or "—"
        fila = [f["marca"], f["modelo"], f["potencia_w"], f["en_almacen"]]
        if vista == "comercial":
            fila.append(f["pendiente_por_llegar"])
        fila += [f["disponible"], f["reservado"], proyectos_txt]
        if vista == "fisica":
            fila.append(f["danados"])
        ws.append(fila)

    for columna in ws.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws.column_dimensions[columna[0].column_letter].width = min(largo + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def teclado_formato(tipo: str) -> InlineKeyboardMarkup:
    """Pregunta cómo quiere verlo la persona, antes de generar nada."""
    return InlineKeyboardMarkup([[
        InlineKeyboardButton("💬 En el chat", callback_data=f"formato:{tipo}:chat"),
        InlineKeyboardButton("📊 En Excel", callback_data=f"formato:{tipo}:excel"),
    ]])


async def construir_reporte_reservas():
    """Lista de reservas activas (no despachadas), con su estado de Odoo y de despacho."""
    pool = await get_pool()
    async with pool.acquire() as conn:
        filas = await conn.fetch(
            "SELECT r.id, r.marca, r.modelo, r.potencia_w, r.cantidad, r.proyecto, "
            "r.estado_odoo, r.estado_despacho, r.fecha_solicitud, u.nombre AS solicitante "
            "FROM reservas r LEFT JOIN usuarios u ON u.telegram_id = r.solicitado_por "
            "WHERE r.estado_despacho != 'despachada' ORDER BY r.fecha_solicitud"
        )

    if not filas:
        return "No hay reservas activas en este momento.", []

    etiqueta_despacho = {"pendiente": "sin despachar", "parcial": "despacho parcial"}
    lineas = ["*Reservas activas:*", ""]
    for f in filas:
        odoo = "✅" if f["estado_odoo"] == "confirmada" else "⏳ Odoo"
        despacho = etiqueta_despacho.get(f["estado_despacho"], f["estado_despacho"])
        lineas.append(
            f"{odoo} #{f['id']} — {f['cantidad']} × {f['marca']} {f['modelo']} {f['potencia_w']}W — "
            f"{f['proyecto']} ({despacho})"
        )

    return "\n".join(lineas), filas


def construir_excel_reservas(filas) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reservas"

    encabezados = [
        "ID", "Marca", "Modelo", "Potencia (W)", "Cantidad", "Proyecto",
        "Solicitado por", "Confirmado en Odoo", "Estado despacho", "Fecha solicitud",
    ]
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True)

    for f in filas:
        ws.append([
            f["id"], f["marca"], f["modelo"], f["potencia_w"], f["cantidad"], f["proyecto"],
            f["solicitante"] or "—", "Sí" if f["estado_odoo"] == "confirmada" else "No",
            f["estado_despacho"], f["fecha_solicitud"].strftime("%Y-%m-%d %H:%M"),
        ])

    for columna in ws.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws.column_dimensions[columna[0].column_letter].width = min(largo + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


async def construir_reporte_movimientos():
    """Últimas entradas (ordenes_compra recibidas) y salidas (despachos)."""
    pool = await get_pool()
    async with pool.acquire() as conn:
        entradas = await conn.fetch(
            "SELECT marca, modelo, potencia_w, cantidad, proveedor, numero_orden, tipo_entrega, "
            "fecha_recepcion FROM ordenes_compra WHERE estado = 'recibida' "
            "ORDER BY fecha_recepcion DESC LIMIT 300"
        )
        salidas = await conn.fetch(
            "SELECT marca, modelo, potencia_w, cantidad_declarada, destino, tipo, origen_compra, fecha "
            "FROM despachos ORDER BY fecha DESC LIMIT 300"
        )

    lineas = ["📥 *Últimas entradas:*", ""]
    if entradas:
        for e in entradas[:10]:
            lineas.append(
                f"• {e['fecha_recepcion']:%d/%m} — {e['cantidad']} × {e['marca']} {e['modelo']} "
                f"{e['potencia_w']}W — {e['proveedor']}"
            )
    else:
        lineas.append("(ninguna)")

    lineas.append("")
    lineas.append("📤 *Últimas salidas:*")
    if salidas:
        for s in salidas[:10]:
            lineas.append(
                f"• {s['fecha']:%d/%m} — {s['cantidad_declarada']} × {s['marca']} {s['modelo']} "
                f"{s['potencia_w']}W → {s['destino']}"
            )
    else:
        lineas.append("(ninguna)")

    return "\n".join(lineas), entradas, salidas


def construir_excel_movimientos(entradas, salidas) -> io.BytesIO:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Entradas"
    ws1.append(["Fecha", "Marca", "Modelo", "Potencia (W)", "Cantidad", "Proveedor", "N° Orden", "Tipo"])
    for celda in ws1[1]:
        celda.font = Font(bold=True)
    for e in entradas:
        ws1.append([
            e["fecha_recepcion"].strftime("%Y-%m-%d %H:%M") if e["fecha_recepcion"] else "",
            e["marca"], e["modelo"], e["potencia_w"], e["cantidad"],
            e["proveedor"], e["numero_orden"] or "—", e["tipo_entrega"] or "—",
        ])
    for columna in ws1.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws1.column_dimensions[columna[0].column_letter].width = min(largo + 2, 40)

    ws2 = wb.create_sheet("Salidas")
    ws2.append(["Fecha", "Marca", "Modelo", "Potencia (W)", "Cantidad", "Destino", "Tipo", "Origen compra"])
    for celda in ws2[1]:
        celda.font = Font(bold=True)
    for s in salidas:
        ws2.append([
            s["fecha"].strftime("%Y-%m-%d %H:%M"), s["marca"], s["modelo"], s["potencia_w"],
            s["cantidad_declarada"], s["destino"], s["tipo"], s["origen_compra"],
        ])
    for columna in ws2.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in columna)
        ws2.column_dimensions[columna[0].column_letter].width = min(largo + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
